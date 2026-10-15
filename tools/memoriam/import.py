"""Import the supplied MSOA workbook without inventing dates or biographies."""
import sys, json, datetime, hashlib, re, unicodedata
from pathlib import Path
import openpyxl
source=Path(sys.argv[1]); root=Path(__file__).resolve().parents[2]
sheet=openpyxl.load_workbook(source,data_only=True).active
entries=[]; seen={}; slugs=set()
for row,values in enumerate(sheet.values,1):
    if row==1: continue
    last,first,date=values[:3]
    if not first or not last: raise ValueError(f'Missing name at row {row}')
    first=' '.join(str(first).split());last=' '.join(str(last).split())
    exact=isinstance(date,datetime.datetime)
    passing=date.strftime('%B %-d, %Y') if exact else str(date).strip() if date is not None else ''
    key=(first,last,passing)
    if key in seen:
        seen[key]['sourceRows'].append(row);continue
    base=unicodedata.normalize('NFKD',first.split()[0]+'-'+last).encode('ascii','ignore').decode().lower()
    slug=re.sub('[^a-z0-9]+','-',base).strip('-')
    if slug in slugs: slug+=f'-{row}'
    slugs.add(slug)
    item={'id':f'msoa-{row}','slug':slug,'firstName':first,'lastName':last,'passing':passing,'dateISO':date.strftime('%Y-%m-%d') if exact else None,'reviewDate':bool(exact and date.year<1950),'sourceRows':[row]}
    seen[key]=item;entries.append(item)
data={'source':'Manned Spaceflight Operations Association (MSOA)','sourceUrl':'https://www.mannedspaceops.org/','sourceUpdated':'May 1, 2026','sourceFile':source.name,'sourceSha256':hashlib.sha256(source.read_bytes()).hexdigest(),'sourceRowCount':sheet.max_row-1,'entries':entries}
(root/'content/memoriam/entries.json').write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')
print(f'{len(entries)} entries; {sheet.max_row-1} source listings; {sum(len(e["sourceRows"])-1 for e in entries)} exact repeats combined.')
